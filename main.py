import csv
import re
import sys
import tkinter
import openpyxl
import win32api
import threading
from pathlib import Path
from docx import Document
from tkinter import messagebox, ttk
from tkinter import *  # from tkinter import Tk for Python 3.x
from tkinter.filedialog import askopenfilename
from tkinter.filedialog import askopenfilenames

encapsulation_char = "$"
data: [[str]] = []
headers: [str] = []
templates: [str] = []
window = Tk()
dataListBox = ""
templateListBox = ""
printFiles = False

encapsulation_char = "$"

def createOutputIfNotExist():
    Path("./output").mkdir(parents=True, exist_ok=True)

def loadData():
    Tk().withdraw()  # we don't want a full GUI, so keep the root window from appearing
    filename = askopenfilename(filetypes=[("CSV Files", "*.csv")])
    global dataListBox
    dataListBox.delete(0, END)

    global headers
    headers.clear()
    global data
    data.clear()

    if filename == "":
        return

    with open(filename, mode="r", newline="", encoding="cp437") as file:
        reader = csv.reader(file, delimiter="─")
        # Iterate through each row in the CSV file
        for row in reader:
            data.append(row[0].split(";"))

    # Take the first row that contains the header as guide for what the columns are
    headers = data.pop(0)
    for header in headers:
        dataListBox.insert(END, "$"+header+"$")
    window.update()

    templatesBtn = Button(text="Load Templates", command=loadTemplates)
    templatesBtn.grid(row=0, column=1)

    global templateListBox
    templateListBox = Listbox(window, selectmode=NONE, width=24)
    templateListBox.grid(row=1, column=1)

    global dataLabel
    dataLabel["text"] = "There are {} users".format(len(data))
    dataLabel.update()


def loadTemplates():
    Tk().withdraw()  # we don't want a full GUI, so keep the root window from appearing
    global templates
    templates = []
    templates = askopenfilenames(
        filetypes=[
            ("Supported Files", ".docx .xlsx"),
            ("Excel Files", "*.xlsx"),
            ("Word Files", "*.docx"),
        ]
    )
    templateListBox.delete(0, END)
    for template in templates:
        templateListBox.insert(END, template.split("/")[-1])
    window.update()

    global templatesLabel
    templatesLabel = Label(text="There are {} templates".format(len(templates)))
    templatesLabel.grid(row=2, column=1)

    global generateBtn
    generateBtn = Button(text="Generate", command=executeUpdate)
    generateBtn.grid(row=0, column=2)


def updateWord(filePath, employe):
    doc = Document(filePath)

    for key in headers:
        for paragraph in doc.paragraphs:
            comparer = encapsulation_char + key + encapsulation_char
            if comparer in paragraph.text:
                paragraph.text = paragraph.text.replace(
                    comparer, employe[headers.index(key)]
                )

    outputFile = "output/" + employe[headers.index("emplNumEmploye")] + ".docx"
    doc.save(outputFile)
    return outputFile


def updateXLSX(filePath, employe):
    wb = openpyxl.load_workbook(filePath)
    # Itterate through each worksheet in the workbook
    for ws in wb.worksheets:
        # Itterate through each row in the worksheet
        for row in ws.iter_rows():
            # Itterate through each cell in the row
            for cell in row:
                # Check if the cell contains any of the headers from the CSV file which are used as rempalcementkeys
                for key in headers:
                    comparer = encapsulation_char + key.lower() + encapsulation_char
                    if str(cell.value).lower().find(comparer) != -1:
                        remove_word = re.compile(re.escape(comparer), re.IGNORECASE)
                        if employe[0] == "112132" and key == "emplNumEmploye":
                            print(key)
                        newValue = employe[headers.index(key)]
                        if key == "emplNaissance":
                            newValue = newValue.split("/")
                            newValue = (
                                newValue[2] + "/" + newValue[1] + "/" + newValue[0]
                            )
                        elif (
                            key == "ipTauxHoraire"
                            or key == "ipTauxLesson"
                            or key == "ipPourcCommission"
                        ):
                            newValue += "$"
                        elif key == "emplPagerCell" and newValue.strip() == "":
                            newValue = employe[headers.index("emplTelephone")]
                        cell.value = remove_word.sub(newValue, comparer)
    outputFile = "output/" + employe[headers.index("emplNumEmploye")] + ".xlsx"
    wb.save(outputFile)
    return outputFile

def process_template(template, employe):
    if template.endswith(".xlsx"):
        updateXLSX(template, employe)
    elif template.endswith(".docx"):
        updateWord(template, employe)
    # elif template.endswith(".pdf"):
    #     outputFile = updatePDF(template, employe)
    else:
        messagebox.showerror("Error", "Invalid file format")

def executeUpdate():
    # create outputfolder if not exist
    createOutputIfNotExist()
    
    def worker(template, employe, done_event, progress_var):
        process_template(template, employe)
        done_event.set()
        progress_var.set(progress_var.get() + 1)

    threads = []
    done_events = []
    total_tasks = len(data) * len(templates)
    
    # Create a progress bar
    progress_var = tkinter.IntVar()
    progress_bar = ttk.Progressbar(window, maximum=total_tasks, variable=progress_var)
    progress_bar.grid(row=3, column=0, columnspan=2, pady=10)

    # Iterate through the list of rows containing each employee's data
    for employe in data:
        for template in templates:
            done_event = threading.Event()
            done_events.append(done_event)
            thread = threading.Thread(target=worker, args=(template, employe, done_event, progress_var))
            threads.append(thread)
            thread.start()

    # Wait for all threads to complete
    for done_event in done_events:
        done_event.wait()

    # Show success message
    messagebox.showinfo("Success", "All templates have been processed successfully")


def main():
    dataBtn = Button(text="Load Data", command=loadData)
    dataBtn.grid(row=0, column=0)

    global dataListBox
    dataListBox = Listbox(window, selectmode=NONE, width=0, height=0)
    dataListBox.grid(row=1, column=0)

    global dataLabel
    dataLabel = Label(text="There are {} users".format(len(data)))
    dataLabel.grid(row=2, column=0)


    c1 = Checkbutton(window, text='Print',variable=printFiles, onvalue=True, offvalue=False)
    c1.grid(row=1, column=2)


    window.mainloop()

    window.protocol("WM_DELETE_WINDOW", sys.exit())


if __name__ == "__main__":
    main()
